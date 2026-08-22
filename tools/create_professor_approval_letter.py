#!/usr/bin/env python3
"""Create the editable Word version of the NEBULA professor approval letter."""

from pathlib import Path
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_run_font(run, size=11, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_body(document, text):
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208
    paragraph.add_run(text)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208
    paragraph.add_run(text)
    return paragraph


def add_header_footer(section):
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_p.add_run("NEBULA  |  RESEARCH PROJECT REVIEW")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("NEBULA | Draft for academic review")
    set_run_font(run, size=8.5, color=MUTED)


def build_document(output):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    configure_styles(document)
    add_header_footer(section)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("NEBULA")
    set_run_font(run, size=12, color=MUTED, bold=True)
    kicker.paragraph_format.space_after = Pt(6)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Request for Academic Approval and Supervision")
    set_run_font(run, size=22, color=NAVY, bold=True)
    title.paragraph_format.space_after = Pt(8)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Trust-Aware Hybrid Retrieval for Explainable and Fresh Engineering Knowledge Search")
    set_run_font(run, size=10.5, color=MUTED, bold=True)
    subtitle.paragraph_format.space_after = Pt(20)

    metadata = document.add_table(rows=4, cols=2)
    metadata.autofit = False
    metadata.columns[0].width = Inches(1.45)
    metadata.columns[1].width = Inches(5.05)
    rows = [
        ("Applicant", "[Your full name]"),
        ("Programme", "[Programme and department]"),
        ("University", "[University name]"),
        ("Date", "[Date]"),
    ]
    for row, (label, value) in zip(metadata.rows, rows):
        row.cells[0].width = Inches(1.45)
        row.cells[1].width = Inches(5.05)
        set_cell_shading(row.cells[0], "F4F6F9")
        label_run = row.cells[0].paragraphs[0].add_run(label)
        set_run_font(label_run, size=9.5, color=NAVY, bold=True)
        value_run = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(value_run, size=9.5, color=MUTED)
    document.add_paragraph().paragraph_format.space_after = Pt(2)

    add_body(document, "To: [Professor’s full name]  |  [Professor’s title and department]  |  [University name]")
    subject = document.add_paragraph()
    subject.paragraph_format.space_after = Pt(10)
    run = subject.add_run("Subject: Request for approval and supervision of the NEBULA research-based software project")
    set_run_font(run, size=11, color=NAVY, bold=True)
    add_body(document, "Dear Professor [Surname],")
    add_body(document, "I am writing to request your academic review, approval, and—if you consider it suitable—supervision of my research-based software project, NEBULA. The project is an open-source, reproducible search system designed for engineering knowledge bases, where a useful result should be not only relevant but also current, authoritative, explainable, and resilient to distributed-system failures.")

    document.add_heading("Project overview", level=1)
    add_body(document, "Conventional search systems commonly optimise lexical or semantic relevance in isolation. In engineering environments, this can produce results that are technically related but stale, weakly sourced, difficult to verify, or unavailable during a shard or service failure. NEBULA investigates whether hybrid retrieval combined with explicit trust signals can improve the quality and verifiability of engineering knowledge search.")
    quote = document.add_paragraph()
    quote.paragraph_format.left_indent = Inches(0.25)
    quote.paragraph_format.right_indent = Inches(0.25)
    quote.paragraph_format.space_after = Pt(10)
    run = quote.add_run("Research question: Does adding freshness, source-authority, graph, and evidence signals to hybrid lexical-semantic retrieval improve retrieval quality and user verifiability compared with lexical-only, semantic-only, and conventional hybrid baselines?")
    set_run_font(run, size=11, color=NAVY, italic=True)
    add_body(document, "The project is innovative as an integrated research and engineering system rather than as a claim that any one individual algorithm is new. It connects retrieval evaluation, trust-aware ranking, graph authority, evidence presentation, reproducibility controls, and failure-aware distributed search in one inspectable implementation.")

    document.add_heading("Current implementation status", level=1)
    for item in [
        "BM25, semantic hashing, exact vector, HNSW, and hybrid retrieval baselines.",
        "Freshness, source-authority, graph/PageRank, and trust-aware ranking signals.",
        "Evidence-oriented result explanations and deterministic evaluation reports.",
        "In-process distributed coordination using consistent hashing and fan-out result merging.",
        "Replicated shard storage, bounded retry behaviour, health monitoring, circuit breaking, and partial-result reporting.",
        "Docker deployment, Prometheus-compatible metrics, automated tests, and continuous integration.",
        "Versioned synthetic corpus, query labels, benchmark artifacts, experiment manifests, and research-report generation.",
        "Annotation validation, inter-annotator agreement analysis, adjudication support, and a release-quality gate.",
        "A frozen development/held-out query split to reduce evaluation leakage during future tuning.",
    ]:
        add_bullet(document, item)
    add_body(document, "The current evaluation materials are deliberately described as a synthetic engineering fixture and regression baseline. I will not present them as human-subject evidence or as proof of generalisation. The repository records limitations, reproducibility requirements, and the distinction between measured implementation behaviour and future research claims.")

    document.add_heading("Proposed academic study", level=1)
    for item in [
        "Freeze the corpus, query set, labels, configurations, and evaluation protocol.",
        "Compare lexical-only, semantic-only, hybrid, and trust-aware variants using Precision@k, Recall@k, MRR, and NDCG.",
        "Measure latency, HNSW recall, index size, and failure-recovery behaviour.",
        "Run ablation studies for freshness, authority, graph, and evidence signals.",
        "If approved and feasible, conduct a small user evaluation of result verification speed and confidence using informed consent, anonymised records, and an ethics-approved protocol.",
        "Analyse failure cases, annotator agreement, limitations, and threats to validity before preparing a dissertation chapter or publication manuscript.",
    ]:
        add_number(document, item)
    add_body(document, "The principal hypotheses are that hybrid retrieval will improve ranking quality over BM25 alone, trust signals will reduce stale or weakly sourced results, and explicit evidence will improve users’ ability to verify a result. These hypotheses will remain open until the relevant experiments are completed.")

    document.add_heading("Ethics, privacy, and scope", level=1)
    add_body(document, "No human-subject study, participant recruitment, personal data collection, or publication submission will begin without the university’s required approval and guidance. The current repository uses public-safe synthetic data only. Any future human evaluation will use the university’s prescribed consent, storage, anonymisation, withdrawal, and data-retention procedures. I will also seek guidance on whether the proposed user evaluation requires ethics review or an exemption determination.")

    document.add_heading("Request for your review", level=1)
    for item in [
        "Review whether NEBULA is suitable as a university research project.",
        "Advise on the research question, evaluation design, and expected academic contribution.",
        "Confirm whether you would be willing to supervise or recommend an appropriate supervisor.",
        "Advise on ethics approval, participant recruitment, and data-management requirements.",
        "Suggest any changes needed before the formal study phase begins.",
    ]:
        add_bullet(document, item)
    add_body(document, "Repository: https://github.com/Bishrav/NEBULA")
    add_body(document, "I have attached or can provide the system architecture, research plan, evaluation protocol, current benchmark reports, and a demonstration of the running prototype. I would welcome the opportunity to discuss the project and revise its scope in accordance with your guidance.")
    add_body(document, "Thank you for considering this request.")
    add_body(document, "Yours sincerely,")
    add_body(document, "[Your full name]\n[Student ID, if applicable]\n[Programme / Department]\n[University name]\n[Email address]\n[Phone number, optional]")

    document.add_heading("Suggested attachments", level=1)
    for item in [
        "NEBULA project repository and README",
        "System architecture and deployment documentation",
        "NEBULA research plan and study-registration draft",
        "Evaluation and annotation protocol",
        "Current synthetic benchmark reports and experiment manifest",
        "Demonstration URL or local run instructions",
    ]:
        add_bullet(document, item)

    properties = document.core_properties
    properties.title = "NEBULA — Request for Academic Approval and Supervision"
    properties.subject = "Research project approval and supervision request"
    properties.author = "NEBULA Research Team"
    properties.keywords = "NEBULA, information retrieval, research proposal, academic supervision"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: create_professor_approval_letter.py OUTPUT.docx")
    build_document(Path(sys.argv[1]).resolve())
    print(sys.argv[1])
